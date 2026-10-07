"""Build the 画面詳細設計書 template.

Usage: python3 build_template.py <out.xlsx> [layout.json]

layout.json is optional; every key is optional:
  item_columns   [[header, width], ...]  the 項目定義 columns, in order (at least 10)
  item_rows      blank rows of 項目定義
  table_rows     {table title: blank rows} for the tables below 項目定義
  revision_rows  blank rows of 改訂履歴
Metadata and the tables below 項目定義 are laid out by position, so a removed, added or renamed
column moves them instead of breaking the build.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

if len(sys.argv) not in (2, 3):
    raise SystemExit(__doc__)
OUT = sys.argv[1]
LAYOUT = json.load(open(sys.argv[2], encoding="utf-8")) if len(sys.argv) == 3 else {}
wb = Workbook()

thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="1F4E78")
LABEL = PatternFill("solid", fgColor="DDEBF7")
SECTION = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def box(ws, r1, c1, r2, c2, value=None, fill=None, font=None, align=WRAP):
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(r1, c1)
    if value is not None:
        cell.value = value
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c)
            x.border = BOX
            if fill:
                x.fill = fill
    if font:
        cell.font = font
    cell.alignment = align
    return cell


def section(ws, row, c1, title, c2):
    box(ws, row, c1, row, c2, title, fill=SECTION, font=Font(bold=True, size=12),
        align=Alignment(vertical="center"))
    ws.row_dimensions[row].height = 22


def table(ws, row, columns, rows):
    """columns: (header, first_col, last_col). Returns the row after the table."""
    for header, c1, c2 in columns:
        box(ws, row, c1, row, c2, header, fill=HEAD, font=WHITE_B, align=CENTER)
    ws.row_dimensions[row].height = 30
    for r in range(row + 1, row + 1 + rows):
        for _, c1, c2 in columns:
            box(ws, r, c1, r, c2)
    return row + 1 + rows


def dv_list(ws, values, ref, strict):
    # strict=False: suggestions only, any other value is accepted silently.
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(ref)


# ── 改訂履歴 ─────────────────────────────────────────────────────
hist = wb.active
hist.title = "改訂履歴"
hist.sheet_view.showGridLines = False
for col, w in zip("ABCDEF", (2, 8, 14, 20, 30, 70)):
    hist.column_dimensions[col].width = w
hist["B1"] = "改訂履歴"
hist["B1"].font = Font(bold=True, size=16, color="1F4E78")
end = table(hist, 3, [("版", 2, 2), ("改訂日", 3, 3), ("改訂者", 4, 4), ("改訂箇所", 5, 5),
                      ("改訂内容", 6, 6)], LAYOUT.get("revision_rows", 20))
for r in range(4, end):
    hist.cell(r, 3).number_format = "yyyy/mm/dd"

# ── 画面詳細設計 ─────────────────────────────────────────────────
ws = wb.create_sheet("画面詳細設計")
wb.active = 1
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 90
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 85  # numbered image, about 600 px
ws.column_dimensions["C"].width = 2

DEFAULT_COLS = [
    ("No", 6), ("項目名", 20), ("画面上表示ラベル", 20), ("項目種別", 14), ("説明", 40), ("操作", 14),
    ("遷移先", 18), ("操作時の動作", 36), ("データ型", 12), ("必須", 9), ("形式", 14), ("最大桁数", 9),
    ("最小桁数", 9), ("初期値", 14), ("入力チェック", 30), ("テーブル名", 18), ("カラム名", 18),
    ("データベース備考", 40),
]
DESC_COLS = [tuple(x) for x in LAYOUT.get("item_columns", DEFAULT_COLS)]
if len(DESC_COLS) < 10 or len({h for h, _ in DESC_COLS}) != len(DESC_COLS):
    raise SystemExit("item_columns needs at least 10 columns with distinct headers")
FIRST = 4  # column D
LAST = FIRST + len(DESC_COLS) - 1
for i, (_, w) in enumerate(DESC_COLS):
    ws.column_dimensions[L(FIRST + i)].width = w
col = {h: FIRST + i for i, (h, _) in enumerate(DESC_COLS)}
DEFAULT_INDEX = {h: i for i, (h, _) in enumerate(DEFAULT_COLS)}


def c(h):
    """Column of a default header, mapped by position onto the actual column list."""
    i = DEFAULT_INDEX[h]
    return FIRST + round(i * (len(DESC_COLS) - 1) / (len(DEFAULT_COLS) - 1))


def spans(pairs):
    """Map (first_header, last_header) pairs of one row onto non overlapping column spans."""
    out, prev = [], FIRST - 1
    for k, (a, b) in enumerate(pairs):
        room = LAST - (len(pairs) - 1 - k)
        c1 = min(max(c(a), prev + 1), room)
        c2 = min(max(c(b), c1), room)
        out.append((c1, c2))
        prev = c2
    return out

ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=min(FIRST + 3, c("説明") - 1))
ws["B1"] = "画面詳細設計書"
ws["B1"].font = Font(bold=True, size=16, color="1F4E78")
ws.row_dimensions[1].height = 30
legend = box(ws, 1, c("説明"), 1, LAST,
             "凡例　赤字：既存のDBに存在しない新規追加の項目。データベース備考に「新規追加」と追加理由を記載する。"
             "空欄：該当なし。未確定の内容は空欄にせず、未決事項に登録してNoを記載する。",
             align=Alignment(wrap_text=True, vertical="center"))
legend.font = Font(size=9, color="595959")
for cc in range(c("説明"), LAST + 1):
    ws.cell(1, cc).border = Border()

# Metadata in three rows: 2 = name, status, design link; 3 = overview; 4 = dates and people.
lab = lambda r, col, text: box(ws, r, col, r, col, text, fill=LABEL, font=BOLD, align=Alignment(vertical="center", wrap_text=True))
DATE = "yyyy/mm/dd"
lab(2, 2, "画面名")
lab(3, 2, "画面概要")
lab(4, 2, "作成日・作成者")
# (first header, last header, label text or None for a value cell, number format)
ROW2 = [("No", "説明", None, None), ("操作", "操作", "ステータス", None), ("遷移先", "遷移先", None, None),
        ("操作時の動作", "操作時の動作", "デザイン参照", None), ("データ型", "データベース備考", None, None)]
ROW4 = [("No", "項目名", None, DATE), ("画面上表示ラベル", "画面上表示ラベル", None, None),
        ("項目種別", "項目種別", "最終更新", None), ("説明", "説明", None, DATE), ("操作", "操作", None, None),
        ("遷移先", "遷移先", "レビュー", None), ("操作時の動作", "操作時の動作", None, DATE),
        ("データ型", "必須", None, None), ("形式", "形式", "デザイン確認", None),
        ("最大桁数", "最小桁数", None, DATE)]
for r, cells in ((2, ROW2), (4, ROW4)):
    for (c1, c2), (_, _, label, fmt) in zip(spans([(a, b) for a, b, _, _ in cells]), cells):
        if label:
            lab(r, c1, label)
        else:
            cell = box(ws, r, c1, r, c2, align=Alignment(vertical="center", wrap_text=True))
            if fmt:
                cell.number_format = fmt
status_col = spans([(a, b) for a, b, _, _ in ROW2])[2][0]
ws.cell(2, status_col).value = "作成中"
dv_list(ws, ["作成中", "確定"], f"{L(status_col)}2", True)
box(ws, 3, FIRST, 3, LAST)
ws.row_dimensions[3].height = 45

# ── 画面イメージ (left) and 項目定義 (right) ─────────────────────
TOP = 6
section(ws, TOP, 2, "画面イメージ", 2)
section(ws, TOP, FIRST, "項目定義", LAST)
desc_head = TOP + 1
desc_end = table(ws, desc_head, [(h, n, n) for h, n in col.items()], LAYOUT.get("item_rows", 40))
box(ws, desc_head, 2, desc_end - 1, 2, "番号付き画面画像を貼り付ける",
    align=Alignment(horizontal="center", vertical="center", wrap_text=True))
ws.cell(desc_head, 2).font = Font(italic=True, color="808080")
rng = lambda h, a=desc_head + 1, b=desc_end - 1: f"{L(col[h])}{a}:{L(col[h])}{b}"
ITEM_LISTS = [("項目種別", ["ボタン", "チェックボックス", "ラジオボタン", "ドロップダウン", "テキスト入力", "テキストエリア",
             "日付選択", "ファイル・画像", "動画", "ページネーション", "ポップアップ", "ラベル", "その他"]),
              ("操作", ["クリック時", "ホバー時", "キー操作時", "一定時間後"]),
              ("データ型", ["文字列", "整数", "小数", "日付", "日時", "真偽値", "配列", "ファイル", "なし"]),
              ("必須", ["必須", "任意"])]
# Suggestion lists and number checks apply only to the columns the layout keeps under their default name.
for h, values in ITEM_LISTS:
    if h in col:
        dv_list(ws, values, rng(h), False)
for h in (x for x in ("最大桁数", "最小桁数") if x in col):
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng(h))

# ── Tables below 項目定義 ────────────────────────────────────────
# (title, columns as (header, first_name, last_name), rows)
TABLES = [
    ("表示・業務ルール", [("ルールID", "No", "項目名"), ("対象項目No", "画面上表示ラベル", "画面上表示ラベル"),
                         ("ルール", "項目種別", "データベース備考")], 15),
    ("状態定義", [("対象項目No", "No", "項目名"), ("状態", "画面上表示ラベル", "項目種別"),
                  ("条件", "説明", "遷移先"), ("表示・動作", "操作時の動作", "データベース備考")], 15),
    ("エラー処理", [("No", "No", "No"), ("対象項目No", "項目名", "画面上表示ラベル"),
                    ("エラーケース", "項目種別", "説明"), ("発生条件", "操作", "操作時の動作"),
                    ("システムの動作", "データ型", "入力チェック"), ("メッセージID", "テーブル名", "テーブル名"),
                    ("備考", "カラム名", "データベース備考")], 10),
    ("メッセージ一覧", [("メッセージID", "No", "画面上表示ラベル"), ("種別", "項目種別", "項目種別"),
                        ("表示メッセージ", "説明", "初期値"), ("表示方法", "入力チェック", "入力チェック"),
                        ("備考", "テーブル名", "データベース備考")], 15),
    ("未決事項", [("No", "No", "No"), ("対象項目No", "項目名", "画面上表示ラベル"),
                  ("質問", "項目種別", "遷移先"), ("決定内容", "操作時の動作", "入力チェック"),
                  ("決定者", "テーブル名", "テーブル名"), ("決定日", "カラム名", "カラム名"),
                  ("状態", "データベース備考", "データベース備考")], 15),
]
row = desc_end + 1
anchors = {}
for title, cols, n in TABLES:
    n = LAYOUT.get("table_rows", {}).get(title, n)
    section(ws, row, FIRST, title, LAST)
    head = row + 1
    mapped = spans([(a, b) for _, a, b in cols])
    mapped[-1] = (mapped[-1][0], LAST)
    end = table(ws, head, [(h, c1, c2) for (h, _, _), (c1, c2) in zip(cols, mapped)], n)
    anchors[title] = (head, end - 1, {h: c1 for (h, _, _), (c1, _) in zip(cols, mapped)})
    row = end + 1

h0, h1, cm = anchors["メッセージ一覧"]
dv_list(ws, ["エラー", "警告", "情報", "確認", "完了"], f"{L(cm['種別'])}{h0 + 1}:{L(cm['種別'])}{h1}", True)
dv_list(ws, ["インライン", "トースト", "モーダル", "バナー", "アラート"],
        f"{L(cm['表示方法'])}{h0 + 1}:{L(cm['表示方法'])}{h1}", False)
h0, h1, cm = anchors["未決事項"]
dv_list(ws, ["未回答", "回答済み"], f"{L(cm['状態'])}{h0 + 1}:{L(cm['状態'])}{h1}", True)
for r in range(h0 + 1, h1 + 1):
    ws.cell(r, cm["決定日"]).number_format = "yyyy/mm/dd"

ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print(OUT, "desc", desc_head, desc_end - 1, {k: (v[0], v[1]) for k, v in anchors.items()})
